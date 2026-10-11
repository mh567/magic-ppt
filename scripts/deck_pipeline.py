#!/usr/bin/env python3
"""Compile independent prompts; register, review, assemble and verify page images.

Image generation remains with the host tool. Never invent backend calls here.
Paths in deck.json are relative to that JSON file, not the caller's cwd.
"""
import argparse
import copy
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import sys
from pathlib import Path
from datetime import datetime, timezone
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PNS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
REQUIRED_CHECKS = {'text', 'numbers', 'diagram', 'master', 'crop'}
MODES = {'locked-template', 'full-page-image'}
KINDS = {'cover', 'body', 'section', 'closing'}


def dump(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def digest_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def digest(data):
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def resolve(planpath, value):
    p = Path(value)
    return p.resolve() if p.is_absolute() else (Path(planpath).resolve().parent / p).resolve()


def load_plan(path):
    path = Path(path).resolve()
    data = json.loads(path.read_text(encoding='utf-8'))
    data.setdefault('image_provider', 'chatgpt-built-in')
    if data['image_provider'] not in {'chatgpt-built-in', 'host-tool', 'test-fixture'}:
        raise ValueError('Unsupported image_provider')
    if data['image_provider'] == 'host-tool' and not str(data.get('image_tool', '')).strip():
        raise ValueError('host-tool requires image_tool provenance')
    data.setdefault('requested_image_model', 'ChatGPT current built-in image tool'
                    if data['image_provider'] == 'chatgpt-built-in' else 'Current configured host image tool')
    data.setdefault('require_exact_model', False)
    if data.get('schema_version') != 1:
        raise ValueError('schema_version must be 1')
    if data.get('mode') not in MODES:
        raise ValueError('mode must be locked-template or full-page-image')
    if not isinstance(data.get('title'), str) or not data['title'].strip():
        raise ValueError('title is required')
    pages = data.get('slides')
    if not isinstance(pages, list) or not pages:
        raise ValueError('slides must be a non-empty list')
    for i, s in enumerate(pages, 1):
        if s.get('id') != f'P{i:03d}' or s.get('order') != i:
            raise ValueError(f'slide {i}: id/order must be P{i:03d}/{i}')
        if s.get('kind') not in KINDS:
            raise ValueError(f"{s['id']}: unsupported kind")
        for field in ('title', 'speaker_notes', 'layout'):
            if not isinstance(s.get(field), str) or not s[field].strip():
                raise ValueError(f"{s['id']}: {field} is required")
        if not isinstance(s.get('on_slide'), list) or any(not isinstance(x, str) for x in s['on_slide']):
            raise ValueError(f"{s['id']}: on_slide must be a list of exact strings")
        if s['kind'] == 'body' and not s['on_slide']:
            raise ValueError(f"{s['id']}: body content is empty")
        if not isinstance(s.get('sources', []), list):
            raise ValueError(f"{s['id']}: sources must be a list")
    _, bound = template(data, path)
    if bound:
        for s in pages:
            master = bound['masters'].get(s['kind'], {})
            if master.get('navigation_count') and len(s['on_slide']) != master['navigation_count']:
                raise ValueError(s['id'] + ': section navigation must match original three-row template; use body for other structures')
    return path, data


def template(data, planpath, required=False):
    candidate = data.get('template_profile')
    profile_path = resolve(planpath, candidate) if candidate else ROOT / 'assets/master/profile.json'
    if not profile_path.exists():
        if required:
            raise ValueError('Original template not bound. Run template_tools.py bind once; style JSON is not an original PPT.')
        return None, None
    profile = json.loads(profile_path.read_text(encoding='utf-8'))
    src = (profile_path.parent / profile['source_pptx']).resolve()
    if not src.is_file() or digest_file(src) != profile.get('source_sha256'):
        raise ValueError('Template source missing or changed; bind and verify it again')
    return profile_path, profile


def page_digest(data, slide, planpath):
    profile_path, profile = template(data, planpath)
    return digest({'mode': data['mode'], 'page': {k: v for k, v in slide.items() if k not in ('generation', 'qa')},
                   'model_request': data.get('requested_image_model'),
                   'model_required': data.get('require_exact_model', False),
                   'image_provider': data.get('image_provider'),
                   'image_tool': data.get('image_tool'),
                   'profile': profile,
                   'compiled_prompt_sha256': hashlib.sha256(prompt_for(data, slide, planpath).encode('utf-8')).hexdigest(),
                   'style_sha': digest_file(ROOT / 'assets/master-style.json')})


def prompt_for(data, s, planpath):
    profile_path, profile = template(data, planpath)
    style = json.loads((ROOT / 'assets/master-style.json').read_text(encoding='utf-8'))
    header = f"[独立完整逐页提示词 | {s['id']} | {s['kind'].upper()} MASTER]\n"
    body = ['生成一张完整的中文企业演示文稿页面，16:9，1920×1080 构图；清晰可读，无聊天框、无外部边框。',
            '科研院所/国企技术汇报；母版优先级高于内容布局。不得改为赛博仪表盘、霓虹、暗色玻璃卡片或大面积青紫色。',
            '母版固定规则：' + json.dumps(style['masters'][s['kind']], ensure_ascii=False),
            '配色与文字：' + json.dumps(style['colors'], ensure_ascii=False)]
    if profile and s['kind'] in profile['masters']:
        m = profile['masters'][s['kind']]
        body += ['真实模板参考和以下实测坐标优先于上述历史约值；有冲突时服从真实模板。',
                 '原模板实际文字框归一化坐标：' + json.dumps(m.get('role_boxes', {}),ensure_ascii=False)]
        if data['mode'] == 'locked-template' and s['kind'] == 'body':
            body += ['使用提供的真实模板参考。固定背景、Logo、分隔线、副标题条、水印、页码签保持参考位置。',
                     '该页最终使用原生模板固定区；当前图是完整视觉稿，正文将按以下坐标无拉伸裁出。',
                     '正文内容的全部文字、线条和模块仅位于归一化矩形 x,y,w,h：' + str(m.get('content_box')),
                     '正文矩形以外不得放正文内容。完整保留内容矩形内部留白，避免边缘截断。']
            if m.get('requires_transparent_content'):
                body += ['[透明正文画布：本条覆盖以上关于重画底板和上屏标题的绘制要求]',
                         '输出完整16:9透明画布，只绘制 on_slide 正文文字、模块和连线。主标题、副标题、页码仅作上下文，不画入图片。',
                         '不要绘制背景、渐变、Logo、水印、标题条、页码签或透明棋盘格；这些由原生模板提供。',
                         '模块间、画布四周和以下水印保留区必须真实透明，正文排版避开这些矩形：' + str(m.get('transparent_boxes', []))]
        elif data['mode'] == 'locked-template':
            body += ['使用真实模板参考，保持固定背景。仅替换原来的标题/副标题/落款/导航/结束文字。',
                     '封面、章节和结束页最终保留原生模板文字与背景，不新增正文装饰。']
        else:
            body += ['使用真实模板图片作为底板参考；生成视觉近似稿，不宣称原生母版或像素完全相同。']
    else:
        body += ['当前只提供已知风格参数，未提供真实原模板，不能宣称完全复刻。']
    body += ['[精确上屏文字：逐项保留，不改数字、符号、单位或专有名称，不添加其它文案]',
             '主标题：' + s['title'], '副标题：' + s.get('subtitle', '')]
    if s['kind'] == 'body':
        body.append('右下页码：' + str(s['order']))
    body += [f'{i}. {text}' for i, text in enumerate(s['on_slide'], 1)]
    body += ['[正文布局]', s['layout'], '[禁止项]',
             '禁止画入讲稿、来源清单、P001文件名、prompt标记、模型型号标签（除非已列入上屏文字）。',
             '禁止自行增加图标、图例、装饰节点、数据、组织名称或时间。复杂结构优先对齐、留白和细线。']
    return header + '\n'.join(body) + '\n'


def prepare(args):
    path, d = load_plan(args.plan)
    profile_path, profile = template(d, path)
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    jobs, md = [], [f"# {d['title']}：逐页提示词与口播\n"]
    for s in d['slides']:
        p = out / 'prompts' / (s['id'] + '.txt')
        p.parent.mkdir(exist_ok=True)
        prompt = prompt_for(d, s, path)
        p.write_text(prompt, encoding='utf-8')
        _, m = (None, profile['masters'].get(s['kind'])) if profile else (None, None)
        refs = []
        if m and m.get('reference_image'):
            rp = (profile_path.parent / m['reference_image']).resolve()
            if rp.is_file():
                refs.append(str(rp))
        qa = s.get('qa', {})
        reusable = False
        if s.get('generation') and qa.get('status') == 'passed':
            try:
                check_page(d, s, path)
                reusable = True
            except ValueError:
                pass
        jobs.append({'id': s['id'], 'prompt_file': str(p), 'references': refs,
                     'requested_image_model': d['requested_image_model'],
                     'require_exact_model': d.get('require_exact_model', False),
                     'image_provider': d['image_provider'],
                     'image_tool': d.get('image_tool'),
                     'transparent_background': bool(d['mode'] == 'locked-template' and m and m.get('requires_transparent_content')),
                     'status': 'reuse' if reusable else 'pending',
                     'content_hash': page_digest(d, s, path)})
        md += [f"\n## {s['id']} {s['title']}\n\n```text\n{prompt}```\n",
               '### 讲师口播\n\n' + s['speaker_notes'] + '\n',
               '### 来源\n\n' + '\n'.join(map(str, s.get('sources', []))) + '\n']
    (out / 'prompts.md').write_text('\n'.join(md), encoding='utf-8')
    dump(out / 'jobs.json', {'template_bound': bool(profile), 'jobs': jobs})
    print(json.dumps({'pages': len(jobs), 'template_bound': bool(profile), 'out': str(out)}, ensure_ascii=False))


def page_by_id(d, sid):
    for s in d['slides']:
        if s['id'] == sid:
            return s
    raise ValueError('Unknown slide: ' + sid)


def image_info(path, min_width=1920):
    with Image.open(path) as im:
        im.verify()
    with Image.open(path) as im:
        w, h = im.size
    if abs(w / h / (16 / 9) - 1) > 0.002:
        raise ValueError(f'Image must be 16:9, got {w}x{h}; regenerate, do not silently stretch/crop')
    if w < min_width:
        raise ValueError(f'Image width {w} below {min_width}; upscaling does not count as new detail')
    return w, h


def record(args):
    path, d = load_plan(args.plan)
    s = page_by_id(d, args.slide)
    p = Path(args.image).resolve()
    w, h = image_info(p, d.get('min_image_width', 1920))
    if not args.actual_model.strip() or not args.tool.strip():
        raise ValueError('Actual model and tool provenance required')
    attempt = s.get('generation', {}).get('attempt', 0) + 1
    if attempt > d.get('max_attempts_per_slide', 3):
        raise ValueError('Attempt limit reached. Keep completed pages; resolve the failure/budget before further calls.')
    s['generation'] = {'image_path': str(p), 'image_sha256': digest_file(p), 'width': w, 'height': h,
                       'actual_model': args.actual_model, 'tool': args.tool,
                       'attempt': attempt, 'content_hash': page_digest(d, s, path),
                       'recorded_at': datetime.now(timezone.utc).isoformat()}
    s['qa'] = {'status': 'pending'}
    dump(path, d)
    print(s['id'] + ': registered; visual review pending')


def model_check(d, gen):
    if d.get('require_exact_model', False) and gen.get('actual_model') != d.get('requested_image_model'):
        raise ValueError('Requested exact image model is not confirmed by generation record')
    if d.get('image_provider', 'chatgpt-built-in') == 'chatgpt-built-in' and gen.get('tool') not in {'image_gen.imagegen', 'image_gen__imagegen'}:
        raise ValueError('ChatGPT built-in image tool provenance required; do not silently substitute API/CLI or another provider')

    if d.get('image_provider') == 'host-tool' and gen.get('tool') != d.get('image_tool'):
        raise ValueError('Configured host image tool provenance mismatch')

def transparency_check(d, s, path, image):
    if d['mode'] != 'locked-template' or s['kind'] != 'body':
        return
    _, profile = template(d, path)
    m = profile['masters']['body'] if profile else {}
    if not m.get('requires_transparent_content'):
        return
    with Image.open(image) as im:
        if 'A' not in im.getbands():
            raise ValueError('Transparent body PNG required to preserve original watermark')
        alpha = im.getchannel('A')
        if alpha.getextrema()[0] == 255:
            raise ValueError('Opaque body would cover original watermark')
        for x, y, w, h in m.get('transparent_boxes', []):
            crop = alpha.crop((round(x*im.width), round(y*im.height), round((x+w)*im.width), round((y+h)*im.height)))
            hist = crop.histogram()
            if sum(hist[9:]) / sum(hist) > .02:
                raise ValueError('Body content covers protected watermark area; edit with current ChatGPT image tool')


def approve(args):
    path, d = load_plan(args.plan)
    s = page_by_id(d, args.slide)
    gen = s.get('generation')
    if not gen:
        raise ValueError('No generation record')
    p = resolve(path, gen['image_path'])
    image_info(p, d.get('min_image_width', 1920))
    if digest_file(p) != gen['image_sha256'] or page_digest(d, s, path) != gen['content_hash']:
        raise ValueError('Image or page has changed; record current artifact before review')
    model_check(d, gen)
    transparency_check(d, s, path, p)
    checks = set(args.checks.split(','))
    if not REQUIRED_CHECKS <= checks or not args.comment.strip():
        raise ValueError('Review requires text,numbers,diagram,master,crop and a concrete review comment')
    s['qa'] = {'status': 'passed', 'image_sha256': gen['image_sha256'],
               'content_hash': gen['content_hash'], 'checks': sorted(checks),
               'reviewer': args.reviewer, 'comment': args.comment,
               'reviewed_at': datetime.now(timezone.utc).isoformat()}
    dump(path, d)
    print(s['id'] + ': visual review recorded')


def check_page(d, s, path):
    gen, qa = s.get('generation', {}), s.get('qa', {})
    if qa.get('status') != 'passed' or not gen:
        raise ValueError(s['id'] + ': visual QA missing')
    p = resolve(path, gen['image_path'])
    image_info(p, d.get('min_image_width', 1920))
    if digest_file(p) != gen.get('image_sha256') or qa.get('image_sha256') != gen.get('image_sha256'):
        raise ValueError(s['id'] + ': image changed after review')
    current = page_digest(d, s, path)
    if current != gen.get('content_hash') or current != qa.get('content_hash'):
        raise ValueError(s['id'] + ': page/template/prompt changed after review')
    if not REQUIRED_CHECKS <= set(qa.get('checks', [])):
        raise ValueError(s['id'] + ': incomplete visual checklist')
    model_check(d, gen)
    transparency_check(d, s, path, p)
    return p


def role_values(s, master):
    title = s['title']
    if master.get('combine_title_subtitle') and s.get('subtitle'):
        title += '\n' + s['subtitle']
    values = {'title': title, 'subtitle': s.get('subtitle', ''), 'page_number': str(s['order']),
              'footer': s.get('footer', ''), 'closing_text': s['title'], 'navigation': '\n'.join(s['on_slide'])}
    for i, text in enumerate(s['on_slide'], 1):
        values['navigation_' + str(i)] = text
    return values


def replace_text(shape, text):
    """Retain original first paragraph/run formatting instead of setting shape.text."""
    if not shape.has_text_frame:
        raise ValueError(f'Shape {shape.shape_id} is not a text shape')
    tf = shape.text_frame
    first = tf.paragraphs[0]
    pp = copy.deepcopy(first._p.pPr) if first._p.pPr is not None else None
    rp = copy.deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    end = first._p.find('{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr')
    end = copy.deepcopy(end) if end is not None else None
    tf.clear()
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if pp is not None:
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            p._p.insert(0, copy.deepcopy(pp))
        run = p.add_run()
        run.text = line
        if rp is not None:
            run._r.insert(0, copy.deepcopy(rp))
        if end is not None:
            p._p.append(copy.deepcopy(end))


def clone_slide(prs, source):
    """Clone static template slide and remap relationship IDs, including pictures."""
    target = prs.slides.add_slide(source.slide_layout)
    relmap = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        new = (target.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external
               else target.part.rels.get_or_add(rel.reltype, rel.target_part))
        relmap[rel.rId] = new
    root = copy.deepcopy(source._element)
    for e in root.iter():
        for name, value in list(e.attrib.items()):
            if name.startswith('{' + RNS + '}') and value in relmap:
                e.set(name, relmap[value])
    for child in list(target._element):
        target._element.remove(child)
    for child in root:
        # Static template clone deliberately excludes animations and transitions.
        if etree.QName(child).localname in ('cSld', 'clrMapOvr'):
            target._element.append(child)
    target._element.attrib.update(root.attrib)
    # add_slide initializes a cached SlideShapes against its original spTree.
    # The copied cSld has a new tree: invalidate that cached wrapper.
    target.__dict__.pop('shapes', None)
    return target


def remove_originals(prs, count):
    for _ in range(count):
        node = prs.slides._sldIdLst[0]
        prs.part.drop_rel(node.rId)
        prs.slides._sldIdLst.remove(node)


def native_fingerprint(slide, excluded=()):
    """Normalize relationship IDs to actual targets so clone remapping is harmless."""
    def normalize(node):
        node = copy.deepcopy(node)
        for el in node.iter():
            for key, value in list(el.attrib.items()):
                if key.startswith('{' + RNS + '}'):
                    rel = slide.part.rels[value]
                    value = rel.target_ref if rel.is_external else hashlib.sha256(rel.target_part.blob).hexdigest()
                    el.set(key, value)
        return etree.tostring(node, method='c14n').decode()
    shapes = [normalize(s._element) for s in slide.shapes if s.shape_id not in excluded and s.name != 'GeneratedContent']
    bg = slide._element.cSld.bg
    return digest({'shapes': shapes, 'background': normalize(bg) if bg is not None else '',
                   'layout': hashlib.sha256(slide.slide_layout.part.blob).hexdigest(),
                   'attributes': dict(slide._element.attrib)})


def assemble(args):
    path, d = load_plan(args.plan)
    imgs = [check_page(d, s, path) for s in d['slides']]
    if len({digest_file(p) for p in imgs}) != len(imgs):
        raise ValueError('Duplicate page image detected; every page needs its own verified image')
    profile_path, profile = template(d, path, required=d['mode'] == 'locked-template')
    if d['mode'] == 'locked-template':
        if profile.get('visual_status') != 'verified':
            raise ValueError('Template binding must be visually verified before assembly')
        prs = Presentation(str(profile_path.parent / profile['source_pptx']))
        originals = list(prs.slides)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
        originals = []
    for s, img in zip(d['slides'], imgs):
        if d['mode'] == 'locked-template':
            if s['kind'] not in profile['masters']:
                raise ValueError('No bound master for ' + s['kind'])
            m = profile['masters'][s['kind']]
            slide = clone_slide(prs, originals[m['slide_index'] - 1])
            for shape in list(slide.shapes):
                if shape.shape_id in m.get('remove_shape_ids', []):
                    shape._element.getparent().remove(shape._element)
            values = role_values(s, m)
            for role, sid in m.get('role_shapes', {}).items():
                shp = next((x for x in slide.shapes if x.shape_id == sid), None)
                if shp is None:
                    raise ValueError('Bound text shape missing: ' + role)
                replace_text(shp, values[role])
            if s['kind'] == 'body':
                x, y, w, h = m['content_box']
                pic = slide.shapes.add_picture(str(img), int(x * prs.slide_width), int(y * prs.slide_height),
                                               width=int(w * prs.slide_width), height=int(h * prs.slide_height))
                pic.name = 'GeneratedContent'
                pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = x, y, 1-x-w, 1-y-h
                pic._element.nvPicPr.cNvPr.set('descr', s['title'] + '；正文为生成图片')
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            pic = slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
            pic.name = 'GeneratedContent'
            pic._element.nvPicPr.cNvPr.set('descr', s['title'] + '；整页为生成图片')
        slide.notes_slide.notes_text_frame.text = notes_text(s)
    if originals:
        remove_originals(prs, len(originals))
    prs.core_properties.title = d['title']
    prs.core_properties.subject = 'Enterprise presentation; mode=' + d['mode']
    out = Path(args.out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    # Reopen and verify rather than assuming export means a valid deck.
    report = verify_deck(d, path, out)
    dump(out.with_suffix('.qa.json'), report)
    print(json.dumps({'pptx': str(out), 'pages': len(d['slides']), 'verified': report['passed']}, ensure_ascii=False))


def notes_text(s):
    src = '\n'.join(map(str, s.get('sources', [])))
    return s['speaker_notes'] + ('\n\n【来源】\n' + src if src else '')


def verify_deck(d, path, pptx):
    prs = Presentation(str(pptx))
    errors, checks = [], []
    if len(prs.slides) != len(d['slides']):
        errors.append('Page count mismatch')
    if abs(prs.slide_width / prs.slide_height / (16/9) - 1) > 0.002:
        errors.append('Deck is not 16:9')
    _, profile = template(d, path, required=d['mode'] == 'locked-template')
    originals = None
    if profile and d['mode'] == 'locked-template':
        pp, _ = template(d, path, True)
        originals = Presentation(str(pp.parent / profile['source_pptx']))
    for s, slide in zip(d['slides'], prs.slides):
        expected = check_page(d, s, path)
        if slide.notes_slide.notes_text_frame.text != notes_text(s):
            errors.append(s['id'] + ': notes mismatch')
        pics = [x for x in slide.shapes if x.name == 'GeneratedContent']
        need_pic = d['mode'] == 'full-page-image' or s['kind'] == 'body'
        if len(pics) != (1 if need_pic else 0):
            errors.append(s['id'] + ': generated image count mismatch')
        elif need_pic:
            pic = pics[0]
            if hashlib.sha256(pic.image.blob).hexdigest() != digest_file(expected):
                errors.append(s['id'] + ': embedded image mismatch')
            box = profile['masters'][s['kind']]['content_box'] if d['mode'] == 'locked-template' else [0, 0, 1, 1]
            rect = [pic.left/prs.slide_width, pic.top/prs.slide_height, pic.width/prs.slide_width, pic.height/prs.slide_height]
            if any(abs(a-b) > 1e-5 for a, b in zip(rect, box)):
                errors.append(s['id'] + ': image geometry mismatch')
            crop = [pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom]
            target_crop = [box[0], box[1], 1-box[0]-box[2], 1-box[1]-box[3]]
            if any(abs(a-b) > 1e-5 for a, b in zip(crop, target_crop)):
                errors.append(s['id'] + ': image crop mismatch')
        if originals:
            m = profile['masters'][s['kind']]
            exclude = list(m.get('role_shapes', {}).values()) + m.get('remove_shape_ids', [])
            source = originals.slides[m['slide_index'] - 1]
            if native_fingerprint(source, exclude) != native_fingerprint(slide, exclude):
                errors.append(s['id'] + ': original fixed master objects/background/layout changed')
            values = role_values(s, m)
            for role, sid in m.get('role_shapes', {}).items():
                shape = next((x for x in slide.shapes if x.shape_id == sid), None)
                if shape is None or shape.text != values[role]:
                    errors.append(s['id'] + ': original text role mismatch ' + role)
        checks.append({'id': s['id'], 'notes': bool(s['speaker_notes']), 'actual_model': s['generation']['actual_model']})
    report = {'passed': not errors, 'checks': checks, 'errors': errors,
              'visual_render_review': 'required-after-render', 'mode': d['mode'],
              'pptx_sha256': digest_file(pptx)}
    if errors:
        raise ValueError(json.dumps(report, ensure_ascii=False))
    return report


def verify(args):
    path, d = load_plan(args.plan)
    report = verify_deck(d, path, Path(args.pptx).resolve())
    dump(args.out, report)
    print(json.dumps(report, ensure_ascii=False))


def contact_sheet(args):
    paths = sorted(Path(args.images).glob('*.png'), key=lambda p: [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', p.name)])
    if not paths:
        raise ValueError('No rendered PNG pages found')
    cols, tw, th, label = args.columns, 480, 270, 30
    if cols < 1:
        raise ValueError('columns must be positive')
    canvas = Image.new('RGB', (cols*tw, ((len(paths)+cols-1)//cols)*(th+label)), '#e6e9ef')
    draw = ImageDraw.Draw(canvas)
    for i, p in enumerate(paths):
        with Image.open(p) as im:
            im = im.convert('RGB')
            im.thumbnail((tw, th))
            x, y = i % cols * tw, i // cols * (th + label)
            canvas.paste(im, (x, y))
            draw.text((x + 8, y + th + 5), f'{i+1} | {p.stem}', fill='#11298c')
    Path(args.out).resolve().parent.mkdir(parents=True, exist_ok=True)
    canvas.save(args.out, quality=92)
    print(json.dumps({'pages': len(paths), 'contact_sheet': str(Path(args.out).resolve())}))


def render(args):
    pptx = Path(args.pptx).resolve()
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    if any(out.glob('*.png')):
        raise ValueError('Render directory must be free of old PNG pages; use a new directory to avoid stale-page contamination.')
    lo, pp = shutil.which('soffice'), shutil.which('pdftoppm')
    if not lo or not pp:
        raise ValueError('soffice and pdftoppm must be on PATH, or use the host presentation renderer and retain its provenance.')
    # Native Chinese text can render blank with no CJK fonts even when export
    # returns success. Catch this observed failure before accepting PNG output.
    native = Presentation(str(pptx))
    cjk = any(re.search(r'[\u3400-\u9fff]', sh.text) for sl in native.slides for sh in sl.shapes if sh.has_text_frame)
    if cjk and sys.platform.startswith('linux') and shutil.which('fc-list'):
        fonts = subprocess.run(['fc-list',':lang=zh'],capture_output=True,text=True,timeout=20)
        if fonts.returncode == 0 and not fonts.stdout.strip():
            raise ValueError('No Chinese fonts detected. Install the licensed original template fonts or use a renderer with those fonts; do not approve blank/missing Chinese titles.')
    with tempfile.TemporaryDirectory(prefix='ppt-final-render-') as temp:
        t = Path(temp)
        result = subprocess.run([lo, f'-env:UserInstallation={(t / "lo-profile").as_uri()}', '--headless',
                                 '--convert-to','pdf','--outdir',str(t),str(pptx)],capture_output=True,text=True,timeout=180)
        pdf = t / (pptx.stem + '.pdf')
        if result.returncode or not pdf.is_file():
            raise ValueError('PPT render failed: ' + (result.stderr+result.stdout)[-1500:])
        subprocess.run([pp,'-scale-to-x','1920','-scale-to-y','1080','-png',str(pdf),str(out/'slide')],
                       capture_output=True,check=True,timeout=180)
    pages = sorted(out.glob('*.png'))
    count = len(Presentation(str(pptx)).slides)
    if len(pages) != count:
        raise ValueError('Final rendered page count mismatch')
    dump(out / 'render.json', {'pptx_sha256':digest_file(pptx), 'pages':count,
                             'renderer':'LibreOffice + Poppler',
                             'image_sha256':{p.name:digest_file(p) for p in pages}})
    print(json.dumps({'rendered_pages':count,'out':str(out),'visual_review':'pending'}))


def final_review(args):
    path,d = load_plan(args.plan)
    pptx = Path(args.pptx).resolve()
    report = verify_deck(d,path,pptx)
    directory = Path(args.renders).resolve()
    proof = json.loads((directory/'render.json').read_text(encoding='utf-8'))
    if proof['pptx_sha256'] != digest_file(pptx) or proof['pages'] != len(d['slides']):
        raise ValueError('Render belongs to another PPT revision')
    if len(proof['image_sha256']) != len(d['slides']):
        raise ValueError('Render evidence missing pages')
    for name,h in proof['image_sha256'].items():
        fp = directory / name
        if digest_file(fp) != h:
            raise ValueError('Render changed after rendering: ' + name)
    if not args.comment.strip():
        raise ValueError('Concrete final visual-review comment required')
    report['visual_render_review'] = 'passed'
    report['visual_reviewer'] = args.reviewer
    report['visual_comment'] = args.comment
    report['render_evidence'] = proof
    dump(args.out,report)
    print(json.dumps({'passed':True,'final_render_review':'recorded','pages':len(d['slides'])}))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--version', action='version', version='magic-ppt ' + (ROOT / 'VERSION').read_text(encoding='utf-8').strip())
    sp = ap.add_subparsers(dest='command', required=True)
    p = sp.add_parser('prepare'); p.add_argument('--plan', required=True); p.add_argument('--out', required=True); p.set_defaults(func=prepare)
    p = sp.add_parser('record'); p.add_argument('--plan', required=True); p.add_argument('--slide', required=True); p.add_argument('--image', required=True); p.add_argument('--actual-model', required=True); p.add_argument('--tool', required=True); p.set_defaults(func=record)
    p = sp.add_parser('approve'); p.add_argument('--plan', required=True); p.add_argument('--slide', required=True); p.add_argument('--reviewer', required=True); p.add_argument('--checks', required=True); p.add_argument('--comment', required=True); p.set_defaults(func=approve)
    p = sp.add_parser('assemble'); p.add_argument('--plan', required=True); p.add_argument('--out', required=True); p.set_defaults(func=assemble)
    p = sp.add_parser('verify'); p.add_argument('--plan', required=True); p.add_argument('--pptx', required=True); p.add_argument('--out', required=True); p.set_defaults(func=verify)
    p = sp.add_parser('contact-sheet'); p.add_argument('--images', required=True); p.add_argument('--out', required=True); p.add_argument('--columns', type=int, default=3); p.set_defaults(func=contact_sheet)
    p = sp.add_parser('render'); p.add_argument('--pptx', required=True); p.add_argument('--out', required=True); p.set_defaults(func=render)
    p = sp.add_parser('final-review'); p.add_argument('--plan', required=True); p.add_argument('--pptx', required=True); p.add_argument('--renders', required=True); p.add_argument('--reviewer', required=True); p.add_argument('--comment', required=True); p.add_argument('--out', required=True); p.set_defaults(func=final_review)
    args = ap.parse_args()
    try:
        args.func(args)
    except (ValueError, OSError, KeyError, IndexError, subprocess.SubprocessError) as e:
        ap.exit(2, str(e) + '\n')


if __name__ == '__main__':
    main()
