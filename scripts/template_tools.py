#!/usr/bin/env python3
"""Inspect a real template, bind explicit roles, render references and record review."""
import argparse
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from datetime import datetime, timezone
from pptx import Presentation
from deck_pipeline import ROOT, KINDS, dump, digest_file, replace_text


def inspect(args):
    prs = Presentation(args.pptx)
    out = {'slide_size': [prs.slide_width, prs.slide_height], 'slides': []}
    for i, slide in enumerate(prs.slides, 1):
        shapes = []
        for s in slide.shapes:
            item = {'id': s.shape_id, 'name': s.name, 'type': str(s.shape_type),
                    'box': [round(s.left/prs.slide_width, 6), round(s.top/prs.slide_height, 6),
                            round(s.width/prs.slide_width, 6), round(s.height/prs.slide_height, 6)]}
            if s.has_text_frame:
                item['text'] = s.text
            shapes.append(item)
        out['slides'].append({'slide_index': i, 'layout': slide.slide_layout.name, 'shapes': shapes})
    dump(args.out, out)
    print(json.dumps({'slides': len(prs.slides), 'inspection': str(Path(args.out).resolve())}))


def check_config(config, prs):
    if abs(prs.slide_width/prs.slide_height/(16/9)-1) > 0.002:
        raise ValueError('Template is not 16:9. Do not silently alter the original dimensions.')
    masters = config.get('masters', {})
    if not {'cover', 'body'} <= set(masters):
        raise ValueError('At least cover and body masters must be bound')
    allowed_roles = {'title', 'subtitle', 'page_number', 'footer', 'closing_text', 'navigation', 'navigation_1', 'navigation_2', 'navigation_3'}
    for kind, m in masters.items():
        if kind not in KINDS:
            raise ValueError('Unknown master: ' + kind)
        idx = m.get('slide_index')
        if not isinstance(idx, int) or not 1 <= idx <= len(prs.slides):
            raise ValueError(kind + ': slide_index must be 1-based within template')
        slide = prs.slides[idx - 1]
        ids = {s.shape_id: s for s in slide.shapes}
        roles, removal = m.get('role_shapes', {}), m.get('remove_shape_ids', [])
        if not set(roles) <= allowed_roles:
            raise ValueError(kind + ': unsupported role')
        if not isinstance(removal, list) or any(not isinstance(x, int) for x in removal):
            raise ValueError(kind + ': remove_shape_ids must be integer list')
        if len(set(roles.values())) != len(roles):
            raise ValueError(kind + ': one text shape cannot serve multiple roles')
        if set(roles.values()) & set(removal):
            raise ValueError(kind + ': dynamic text role also selected for deletion')
        if kind == 'body':
            required = {'title', 'subtitle'}
            if m.get('page_number_mode') == 'inherited-field':
                field_id = m.get('layout_page_number_shape_id')
                inherited = next((s for s in slide.slide_layout.shapes if s.shape_id == field_id), None)
                if inherited is None or not inherited._element.xpath('.//a:fld[@type="slidenum"]'):
                    raise ValueError('Inherited page number must identify the real layout slidenum field')
            else:
                required.add('page_number')
            if not required <= set(roles):
                raise ValueError('body requires title, subtitle and a native page-number role or inherited-field')
        if kind == 'cover' and 'title' not in roles:
            raise ValueError('cover requires title role')
        for role, sid in roles.items():
            if sid not in ids or not ids[sid].has_text_frame:
                raise ValueError(kind + ': role must identify a top-level original text shape: ' + role)
        if not set(removal) <= set(ids):
            raise ValueError(kind + ': unknown removal shape IDs')
        if kind == 'body':
            box = m.get('content_box')
            if not isinstance(box, list) or len(box) != 4 or any(not isinstance(x, (int,float)) for x in box):
                raise ValueError('body content_box must be [x,y,w,h] normalized numbers')
            x,y,w,h = box
            if min(x,y) < 0 or min(w,h) <= 0 or x+w > 1 or y+h > 1:
                raise ValueError('body content box outside slide')
            for zone in m.get('transparent_boxes', []):
                if len(zone) != 4 or min(zone[:2]) < 0 or min(zone[2:]) <= 0 or zone[0]+zone[2] > 1 or zone[1]+zone[3] > 1:
                    raise ValueError('Transparent watermark box outside slide')
            # Dynamic header/page tab must not be obscured by the opaque content image.
            for role, sid in roles.items():
                s = ids[sid]
                sx,sy,sw,sh = s.left/prs.slide_width,s.top/prs.slide_height,s.width/prs.slide_width,s.height/prs.slide_height
                if max(x,sx) < min(x+w,sx+sw) and max(y,sy) < min(y+h,sy+sh):
                    raise ValueError('body content box overlaps text role: ' + role)
    return masters


def bind(args):
    src, cfg = Path(args.pptx).resolve(), Path(args.config).resolve()
    if not src.is_file():
        raise ValueError('Original PPTX not found')
    config = json.loads(cfg.read_text(encoding='utf-8'))
    prs = Presentation(str(src))
    masters = check_config(config, prs)
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    destination = out / 'source.pptx'
    if destination.exists() and digest_file(destination) != digest_file(src) and not args.replace:
        raise ValueError('A different bound source exists. Preserve it or explicitly use --replace after user requests a template change.')
    if src != destination:
        shutil.copy2(src, destination)
    # Raw and blank previews are rendered separately; the original source stays untouched.
    blank = Presentation(str(destination))
    selected = {m['slide_index']: m for m in masters.values()}
    for idx, m in selected.items():
        slide = blank.slides[idx-1]
        for s in list(slide.shapes):
            if s.shape_id in m.get('remove_shape_ids', []):
                s._element.getparent().remove(s._element)
            elif s.shape_id in m.get('role_shapes', {}).values():
                replace_text(s, '')
        if slide.has_notes_slide:
            slide.notes_slide.notes_text_frame.text = ''
    blank.save(str(out / 'blank-preview.pptx'))
    for kind, m in masters.items():
        source_slide = prs.slides[m['slide_index']-1]
        m['role_boxes'] = {}
        for role, sid in m.get('role_shapes', {}).items():
            sh = next(x for x in source_slide.shapes if x.shape_id == sid)
            m['role_boxes'][role] = [sh.left/prs.slide_width, sh.top/prs.slide_height,
                                     sh.width/prs.slide_width, sh.height/prs.slide_height]
        m['reference_image'] = f'references/{kind}.png'
        m['blank_reference_image'] = f'blank/{kind}.png'
    profile = {'schema_version': 1, 'source_pptx': 'source.pptx',
               'source_sha256': digest_file(destination), 'masters': masters,
               'visual_status': 'pending', 'provenance': config.get('provenance', 'User supplied original PPTX'),
               'bound_at': datetime.now(timezone.utc).isoformat()}
    dump(out / 'profile.json', profile)
    if out == (ROOT / 'assets/master').resolve():
        dump(ROOT / 'assets/template-status.json', {'status': 'bound-awaiting-visual-review',
                                                  'profile': 'master/profile.json'})
    print(json.dumps({'profile': str(out / 'profile.json'), 'visual_status': 'pending'}, ensure_ascii=False))


def executable(name):
    found = shutil.which(name)
    if not found:
        raise ValueError(f'{name} not on PATH. Use the host presentation renderer or add an installed executable to PATH.')
    return found


def render(args):
    profile_path = Path(args.profile).resolve()
    p = json.loads(profile_path.read_text(encoding='utf-8'))
    base = profile_path.parent
    src = base / p['source_pptx']
    if digest_file(src) != p['source_sha256']:
        raise ValueError('Source template changed')
    lo, pp = executable('soffice'), executable('pdftoppm')
    with tempfile.TemporaryDirectory(prefix='ppt-master-render-') as temp:
        t = Path(temp)
        for source_name, output_dir in ((p['source_pptx'], 'references'), ('blank-preview.pptx', 'blank')):
            dest = t / output_dir
            dest.mkdir()
            r = subprocess.run([lo, f'-env:UserInstallation={ (t / (output_dir+"-lo-profile")).as_uri() }',
                                '--headless','--convert-to','pdf','--outdir',str(dest),str(base/source_name)],
                               capture_output=True, text=True, timeout=180)
            pdf = dest / (Path(source_name).stem + '.pdf')
            if r.returncode or not pdf.is_file():
                raise ValueError('Template render failed: ' + (r.stderr + r.stdout)[-1500:])
            (base / output_dir).mkdir(exist_ok=True)
            for kind, m in p['masters'].items():
                idx = m['slide_index']
                prefix = dest / kind
                subprocess.run([pp,'-f',str(idx),'-l',str(idx),'-singlefile','-scale-to-x','1920','-scale-to-y','1080',
                                '-png',str(pdf),str(prefix)],check=True,capture_output=True,timeout=90)
                shutil.copy2(prefix.with_suffix('.png'), base / output_dir / (kind + '.png'))
    print(json.dumps({'references': [str(base / m['reference_image']) for m in p['masters'].values()],
                      'review': 'Compare originals and blank masters visually before mark-verified'}))
    warnings = []
    if shutil.which('fc-list'):
        fonts = subprocess.run(['fc-list', ':lang=zh'], capture_output=True, text=True, timeout=20)
        if fonts.returncode == 0 and not fonts.stdout.strip():
            warnings.append('Chinese fonts missing: reference text renders blank; fixed geometry only was previewed. Re-render with original licensed fonts or a suitable host renderer before acceptance.')
    p['render_warnings'] = warnings
    dump(profile_path, p)


def mark_verified(args):
    path = Path(args.profile).resolve()
    p = json.loads(path.read_text(encoding='utf-8'))
    src = path.parent / p['source_pptx']
    if digest_file(src) != p['source_sha256']:
        raise ValueError('Source changed')
    if p.get('render_warnings'):
        raise ValueError('Template text render unresolved: ' + '; '.join(p['render_warnings']))
    hashes = {}
    for kind,m in p['masters'].items():
        for key in ('reference_image','blank_reference_image'):
            fp = path.parent / m[key]
            if not fp.is_file():
                raise ValueError('Render references first; missing ' + str(fp))
            hashes[kind + ':' + key] = digest_file(fp)
    if not args.comment.strip():
        raise ValueError('Concrete visual review comment required')
    p.update(visual_status='verified', visual_review={'reviewer':args.reviewer,'comment':args.comment,
                                                   'reference_sha256':hashes,
                                                   'reviewed_at':datetime.now(timezone.utc).isoformat()})
    dump(path,p)
    if path == ROOT / 'assets/master/profile.json':
        dump(ROOT / 'assets/template-status.json', {'status':'ready','profile':'master/profile.json'})
    print('Template visual review recorded; persist assets with skill commit/push.')


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sp = ap.add_subparsers(dest='command',required=True)
    p=sp.add_parser('inspect');p.add_argument('--pptx',required=True);p.add_argument('--out',required=True);p.set_defaults(func=inspect)
    p=sp.add_parser('bind');p.add_argument('--pptx',required=True);p.add_argument('--config',required=True);p.add_argument('--out',default=str(ROOT/'assets/master'));p.add_argument('--replace',action='store_true');p.set_defaults(func=bind)
    p=sp.add_parser('render');p.add_argument('--profile',required=True);p.set_defaults(func=render)
    p=sp.add_parser('mark-verified');p.add_argument('--profile',required=True);p.add_argument('--reviewer',required=True);p.add_argument('--comment',required=True);p.set_defaults(func=mark_verified)
    args=ap.parse_args()
    try:args.func(args)
    except (ValueError,OSError,KeyError,subprocess.SubprocessError) as e:ap.exit(2,str(e)+'\n')


if __name__=='__main__':main()
