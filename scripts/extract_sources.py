#!/usr/bin/env python3
"""Extract source text with locators. Never claim OCR or chart understanding."""
import argparse
import csv
import json
from pathlib import Path


def extract(path):
    ext=path.suffix.lower()
    units=[]
    warnings=[]
    if ext in ('.md','.txt','.json','.yaml','.yml'):
        text=path.read_text(encoding='utf-8-sig')
        units=[{'locator':'text','text':text}]
    elif ext in ('.csv','.tsv'):
        with path.open(encoding='utf-8-sig',newline='') as f:
            for i,row in enumerate(csv.reader(f,delimiter='\t' if ext=='.tsv' else ','),1):
                units.append({'locator':f'row {i}','values':row})
    elif ext=='.docx':
        from docx import Document
        doc=Document(path)
        for i,p in enumerate(doc.paragraphs,1):
            if p.text.strip():units.append({'locator':f'paragraph {i}','text':p.text})
        for i,t in enumerate(doc.tables,1):
            units.append({'locator':f'table {i}','rows':[[c.text for c in row.cells] for row in t.rows]})
        warnings.append('Inspect embedded images, text boxes, headers and tracked changes separately.')
    elif ext=='.pptx':
        from pptx import Presentation
        prs=Presentation(path)
        for i,s in enumerate(prs.slides,1):
            texts=[]
            def walk(shapes):
                for sh in shapes:
                    if sh.has_text_frame:texts.append(sh.text)
                    if sh.has_table:texts.extend([' | '.join(c.text for c in row.cells) for row in sh.table.rows])
                    if hasattr(sh,'shapes'):walk(sh.shapes)
            walk(s.shapes)
            units.append({'locator':f'slide {i}','text':'\n'.join(texts),
                          'speaker_notes':s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''})
        warnings.append('Render slides to inspect charts, connectors, images and inherited master text.')
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        formulas=load_workbook(path,data_only=False,read_only=True)
        cached=load_workbook(path,data_only=True,read_only=True)
        for ws in formulas:
            cv=cached[ws.title]
            for row in ws.iter_rows():
                cells=[]
                for c in row:
                    if c.value is not None:
                        cells.append({'cell':c.coordinate,'value':str(c.value),
                                      'cached_value':cv[c.coordinate].value if c.data_type=='f' else None})
                if cells:units.append({'locator':f'{ws.title}!row {row[0].row}','cells':cells})
        formulas.close();cached.close()
        warnings.append('Cached formula values may be missing/stale. Do not treat this as recalculation; inspect charts separately.')
    elif ext=='.pdf':
        from pypdf import PdfReader
        doc=PdfReader(path)
        for i,p in enumerate(doc.pages,1):
            t=p.extract_text() or ''
            units.append({'locator':f'page {i}','text':t})
            if not t.strip():warnings.append(f'page {i}: no text extracted; render/OCR required, not evidence of an empty page.')
    elif ext in ('.doc','.ppt','.xls'):
        warnings.append('Legacy Office format: convert with supported Office/LibreOffice tool, then extract the converted copy.')
    elif ext in ('.png','.jpg','.jpeg','.webp','.tif','.tiff','.bmp'):
        warnings.append('Image source: inspect visually/OCR through host tools. No OCR performed by this script.')
    else:
        warnings.append('Unsupported format: route to host file tools; do not silently omit it.')
    return units,warnings


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('input');ap.add_argument('--out',required=True)
    args=ap.parse_args();p=Path(args.input).resolve()
    files=sorted(x for x in p.rglob('*') if x.is_file() and not any(q.startswith('.') for q in x.relative_to(p).parts)) if p.is_dir() else [p]
    result=[]
    for f in files:
        if f.resolve()==Path(args.out).resolve():continue
        try:
            u,w=extract(f)
            result.append({'file':str(f),'status':'needs_additional_inspection' if w else 'text_extracted','units':u,'warnings':w})
        except Exception as e:
            result.append({'file':str(f),'status':'failed','units':[],'warnings':[str(e)]})
    out=Path(args.out);out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps({'sources':result},ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
    print(json.dumps({'files':len(result),'failed':sum(x['status']=='failed' for x in result),'output':str(out.resolve())}))


if __name__=='__main__':main()
